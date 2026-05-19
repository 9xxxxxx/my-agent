"""数据导出器（Excel/CSV）"""
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows


class DataExporter:
    """数据导出器"""

    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    HEADER_FILL = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
    BORDER = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def export_excel(
        self,
        df: pd.DataFrame,
        output_path: Path,
        title: str = "",
        sheet_name: str = "Sheet1",
    ) -> Path:
        """导出为 Excel 文件"""
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name

        # 如果有标题，添加标题行
        start_row = 1
        if title:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(df.columns))
            title_cell = ws.cell(row=1, column=1, value=title)
            title_cell.font = Font(bold=True, size=14)
            title_cell.alignment = Alignment(horizontal="center")
            start_row = 2

        # 写入数据
        for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), start_row):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                cell.border = self.BORDER

                # 表头样式
                if r_idx == start_row:
                    cell.font = self.HEADER_FONT
                    cell.fill = self.HEADER_FILL
                    cell.alignment = self.HEADER_ALIGNMENT

        # 自动调整列宽
        for column in ws.columns:
            max_length = 0
            column_letter = None
            for cell in column:
                # Skip merged cells (they don't have column_letter)
                if hasattr(cell, 'column_letter'):
                    if column_letter is None:
                        column_letter = cell.column_letter
                    try:
                        if cell.value and len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
            if column_letter:
                adjusted_width = min(max_length + 4, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

        # 冻结首行
        ws.freeze_panes = f"A{start_row + 1}"

        wb.save(str(output_path))
        return output_path

    def export_csv(
        self,
        df: pd.DataFrame,
        output_path: Path,
        encoding: str = "utf-8-sig",
        separator: str = ",",
    ) -> Path:
        """导出为 CSV 文件"""
        df.to_csv(
            output_path,
            index=False,
            encoding=encoding,
            sep=separator,
        )
        return output_path
